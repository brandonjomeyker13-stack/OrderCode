# tests/test_ranking_excel.py
"""
La exportación a Excel.

Lo que se verifica no es que el archivo "se vea bonito" sino que diga
exactamente lo mismo que la pantalla, que sea un .xlsx que Excel pueda abrir,
y que un nombre de taller no pueda convertirse en algo ejecutable dentro del
archivo.
"""
import io
from datetime import date, datetime, timezone
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.excel import ranking_a_excel
from app.models.empresa import Empresa
from app.models.lote import Lote
from app.models.taller import Taller
from app.schemas.asignacion_taller import (
    AsignacionTallerCreate,
    AsignacionTallerRegistrarDevolucion,
)
from app.services.asignacion_taller_service import AsignacionTallerService
from app.services.ranking_service import RankingService

FEBRERO = (date(2026, 2, 1), date(2026, 2, 28))


def _asignar(db, empresa, lote, taller, cantidad):
    return AsignacionTallerService(db).crear(
        empresa.id,
        AsignacionTallerCreate(
            lote_id=lote.id,
            taller_id=taller.id,
            cantidad_asignada=cantidad,
            fecha_asignacion=date(2026, 2, 1),
        ),
    )


def _devolver(db, empresa, asignacion_id, cantidad, dia=date(2026, 2, 5)):
    return AsignacionTallerService(db).registrar_devolucion(
        asignacion_id,
        empresa.id,
        AsignacionTallerRegistrarDevolucion(
            cantidad_devuelta=cantidad,
            fecha_devolucion=datetime(dia.year, dia.month, dia.day, 12, tzinfo=timezone.utc),
        ),
    )


def _abrir(contenido: bytes):
    return load_workbook(io.BytesIO(contenido)).active


def _buscar_fila(hoja, texto: str) -> int:
    for fila in hoja.iter_rows(min_col=1, max_col=1):
        if fila[0].value == texto:
            return fila[0].row
    raise AssertionError(f"no se encontró la fila '{texto}' en el Excel")


def test_el_archivo_es_un_xlsx_que_se_puede_abrir(
    db: Session, empresa: Empresa, lote: Lote, taller: Taller
):
    _asignar(db, empresa, lote, taller, 100)

    contenido, nombre = RankingService(db).exportar_excel(empresa.id, *FEBRERO)

    assert contenido[:2] == b"PK", "un .xlsx es un zip; esto no lo es"
    assert nombre == "talleres-2026-02-01-a-2026-02-28.xlsx"
    assert _abrir(contenido).title == "Talleres"


def test_las_cifras_del_excel_son_las_mismas_de_la_pantalla(
    db: Session, empresa: Empresa, lote: Lote, taller: Taller
):
    """La razón de reusar el servicio en vez de rehacer las consultas: que el
    archivo no pueda contradecir lo que el usuario tiene delante."""
    lote.tipo_producto.precio_unitario = Decimal("1500")
    db.commit()
    asignacion = _asignar(db, empresa, lote, taller, 200)
    _devolver(db, empresa, asignacion.id, 80)

    ranking = RankingService(db).talleres(empresa.id, *FEBRERO)
    contenido, _ = RankingService(db).exportar_excel(empresa.id, *FEBRERO)
    hoja = _abrir(contenido)

    fila = _buscar_fila(hoja, taller.nombre)
    assert hoja.cell(row=fila, column=3).value == ranking.talleres[0].unidades_devueltas == 80
    assert hoja.cell(row=fila, column=4).value == 120000
    assert hoja.cell(row=fila, column=10).value == 120  # pendiente hoy


def test_lleva_una_fila_de_totales(db: Session, empresa: Empresa, lote: Lote, taller: Taller):
    lote.tipo_producto.precio_unitario = Decimal("1000")
    db.commit()
    asignacion = _asignar(db, empresa, lote, taller, 100)
    _devolver(db, empresa, asignacion.id, 60)

    hoja = _abrir(RankingService(db).exportar_excel(empresa.id, *FEBRERO)[0])

    fila = _buscar_fila(hoja, "TOTAL")
    assert hoja.cell(row=fila, column=3).value == 60
    assert hoja.cell(row=fila, column=4).value == 60000


def test_el_periodo_queda_escrito_en_el_archivo(db: Session, empresa: Empresa, taller: Taller):
    """Un reporte de pagos sin el periodo impreso es inservible al mes
    siguiente: nadie recuerda a qué corte correspondía."""
    contenido, _ = RankingService(db).exportar_excel(
        empresa.id, date(2026, 2, 1), date(2026, 2, 15)
    )
    hoja = _abrir(contenido)

    assert hoja.cell(row=2, column=1).value == empresa.nombre
    assert hoja.cell(row=3, column=1).value == "Periodo: 01/02/2026 a 15/02/2026"


def test_sin_entregas_la_celda_de_dias_queda_vacia(
    db: Session, empresa: Empresa, taller: Taller
):
    """Vacía y no en cero: un 0 se leería como "entregó el mismo día"."""
    hoja = _abrir(RankingService(db).exportar_excel(empresa.id, *FEBRERO)[0])

    fila = _buscar_fila(hoja, taller.nombre)
    assert hoja.cell(row=fila, column=8).value is None


def test_avisa_cuando_faltan_precios_por_cargar(
    db: Session, empresa: Empresa, lote: Lote, taller: Taller
):
    """Sin el aviso, un total incompleto se ve igual que un total real y el
    cliente le paga de menos a su taller sin enterarse."""
    asignacion = _asignar(db, empresa, lote, taller, 100)  # producto sin precio
    _devolver(db, empresa, asignacion.id, 40)

    hoja = _abrir(RankingService(db).exportar_excel(empresa.id, *FEBRERO)[0])

    textos = [
        celda[0].value for celda in hoja.iter_rows(min_col=1, max_col=1) if celda[0].value
    ]
    assert any("ATENCIÓN" in t and "40 unidades" in t for t in textos)


def test_sin_faltantes_no_aparece_el_aviso(
    db: Session, empresa: Empresa, lote: Lote, taller: Taller
):
    lote.tipo_producto.precio_unitario = Decimal("1000")
    db.commit()
    asignacion = _asignar(db, empresa, lote, taller, 100)
    _devolver(db, empresa, asignacion.id, 40)

    hoja = _abrir(RankingService(db).exportar_excel(empresa.id, *FEBRERO)[0])

    textos = [
        celda[0].value for celda in hoja.iter_rows(min_col=1, max_col=1) if celda[0].value
    ]
    assert not any("ATENCIÓN" in t for t in textos)


def test_un_nombre_de_taller_no_se_convierte_en_formula(
    db: Session, empresa: Empresa, lote: Lote
):
    """openpyxl marca como FÓRMULA cualquier texto que empiece por '='. Un
    taller llamado así quedaría como fórmula viva dentro del archivo que el
    cliente le manda a su contador."""
    taller = Taller(
        empresa_id=empresa.id,
        nombre="=1+1",
        numero_taller="=A1",
        telefono="3001234567",
    )
    db.add(taller)
    db.commit()
    _asignar(db, empresa, lote, taller, 50)

    hoja = _abrir(RankingService(db).exportar_excel(empresa.id, *FEBRERO)[0])

    fila = _buscar_fila(hoja, "=1+1")
    assert hoja.cell(row=fila, column=1).data_type == "s", "el nombre quedó como fórmula"
    assert hoja.cell(row=fila, column=2).data_type == "s"


def test_no_se_cuelan_talleres_de_otra_empresa(
    db: Session, empresa: Empresa, lote: Lote, taller: Taller
):
    otra = Empresa(nombre="Empresa Vecina", activo=True)
    db.add(otra)
    db.flush()
    ajeno = Taller(
        empresa_id=otra.id, nombre="Taller Ajeno", numero_taller="TA-1", telefono="3009998877"
    )
    db.add(ajeno)
    db.commit()

    hoja = _abrir(RankingService(db).exportar_excel(empresa.id, *FEBRERO)[0])

    textos = [
        celda[0].value for celda in hoja.iter_rows(min_col=1, max_col=1) if celda[0].value
    ]
    assert "Taller Ajeno" not in textos
    assert empresa.nombre in textos


def test_los_talleres_inactivos_quedan_marcados(
    db: Session, empresa: Empresa, lote: Lote, taller: Taller
):
    """Aparecen porque se les puede deber plata, pero quien lea el archivo
    tiene que saber que ese taller ya no está trabajando."""
    asignacion = _asignar(db, empresa, lote, taller, 100)
    _devolver(db, empresa, asignacion.id, 100)
    taller.activo = False
    db.commit()

    hoja = _abrir(RankingService(db).exportar_excel(empresa.id, *FEBRERO)[0])

    _buscar_fila(hoja, f"{taller.nombre} (inactivo)")


def test_la_merma_va_como_porcentaje_de_excel(
    db: Session, empresa: Empresa, lote: Lote, taller: Taller
):
    """Excel guarda los porcentajes como fracción: 0.2 con formato de % se ve
    como 20%. Guardar 20 con ese formato mostraría 2000%."""
    from app.models.enums import TipoIncidencia
    from app.schemas.incidencia import IncidenciaCreate
    from app.services.incidencia_service import IncidenciaService

    asignacion = _asignar(db, empresa, lote, taller, 200)
    IncidenciaService(db).crear(
        empresa.id,
        IncidenciaCreate(
            asignacion_id=asignacion.id,
            tipo=TipoIncidencia.PRODUCTO_DANADO,
            cantidad_afectada=20,
            descripcion="daño",
            fecha=date(2026, 2, 3),
        ),
    )
    _devolver(db, empresa, asignacion.id, 80)

    hoja = _abrir(RankingService(db).exportar_excel(empresa.id, *FEBRERO)[0])

    fila = _buscar_fila(hoja, taller.nombre)
    celda = hoja.cell(row=fila, column=6)
    assert celda.value == 0.2  # 20 de 100 terminadas
    assert "%" in celda.number_format


def test_el_excel_no_se_rompe_sin_talleres(db: Session, empresa: Empresa):
    """Una empresa recién creada que le da al botón el primer día."""
    contenido, _ = RankingService(db).exportar_excel(empresa.id, *FEBRERO)

    hoja = _abrir(contenido)
    fila = _buscar_fila(hoja, "TOTAL")
    assert hoja.cell(row=fila, column=3).value == 0
