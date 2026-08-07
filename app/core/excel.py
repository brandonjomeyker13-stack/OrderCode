# app/core/excel.py
"""
Exportación del ranking de talleres a un archivo .xlsx.

Este archivo es con el que un cliente le paga a sus talleres, así que las
cuentas NO se hacen acá: llegan ya calculadas desde RankingService y este
módulo solo las escribe. Si el Excel calculara por su cuenta, tarde o
temprano diría algo distinto a la pantalla.

Se eligió xlsx y no CSV porque el Excel en español espera ';' como separador
y sin BOM las tildes salen rotas -- un archivo que se ve mal es justo lo que
no se le puede mandar a un cliente.
"""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.fechas import zona
from app.schemas.ranking import RankingTalleres

AZUL = "1F3A5F"
GRIS_FONDO = "F3F4F6"
BLANCO = "FFFFFF"
ROJO = "B91C1C"

_MONEDA = '"$"#,##0.00'
_ENTERO = "#,##0"
_PORCENTAJE = "0.0%"
_DECIMAL = "0.0"

# (título, ancho, formato numérico)
_COLUMNAS = (
    ("Taller", 28, None),
    ("N° taller", 12, None),
    ("Unidades entregadas", 20, _ENTERO),
    ("Total a pagar", 16, _MONEDA),
    ("Unidades en merma", 19, _ENTERO),
    ("% merma", 10, _PORCENTAJE),
    ("Entregas", 10, _ENTERO),
    ("Días promedio", 14, _DECIMAL),
    ("Sin tarifa", 11, _ENTERO),
    ("Pendiente hoy", 14, _ENTERO),
)

_FILA_ENCABEZADO = 6


def _texto(celda, valor: str) -> None:
    """Escribe texto que vino del usuario SIN que Excel lo interprete.

    openpyxl marca como fórmula cualquier cadena que empiece por '=' (está
    comprobado). Un taller llamado "=A1" acabaría siendo una fórmula viva
    dentro del archivo en vez de un nombre. Forzar el tipo a texto lo deja
    como lo que es: un dato.
    """
    celda.value = valor
    celda.data_type = "s"


def _escribir_cabecera(hoja: Worksheet, ranking: RankingTalleres, empresa_nombre: str) -> None:
    titulo = hoja.cell(row=1, column=1)
    _texto(titulo, "Resumen de talleres")
    titulo.font = Font(bold=True, size=14, color=AZUL)

    empresa = hoja.cell(row=2, column=1)
    _texto(empresa, empresa_nombre)
    empresa.font = Font(bold=True, size=11)

    periodo = hoja.cell(row=3, column=1)
    _texto(
        periodo,
        f"Periodo: {ranking.desde.strftime('%d/%m/%Y')} a {ranking.hasta.strftime('%d/%m/%Y')}",
    )

    generado = hoja.cell(row=4, column=1)
    _texto(generado, f"Generado el {datetime.now(zona()).strftime('%d/%m/%Y %H:%M')}")
    generado.font = Font(size=9, italic=True)

    relleno = PatternFill("solid", fgColor=AZUL)
    borde_inferior = Border(bottom=Side(style="thin", color=AZUL))
    for indice, (nombre, ancho, _) in enumerate(_COLUMNAS, start=1):
        celda = hoja.cell(row=_FILA_ENCABEZADO, column=indice)
        _texto(celda, nombre)
        celda.font = Font(bold=True, color=BLANCO)
        celda.fill = relleno
        celda.border = borde_inferior
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    hoja.row_dimensions[_FILA_ENCABEZADO].height = 30


def _escribir_filas(hoja: Worksheet, ranking: RankingTalleres) -> int:
    """Devuelve el número de la última fila escrita."""
    fila = _FILA_ENCABEZADO
    for taller in ranking.talleres:
        fila += 1
        nombre = taller.nombre if taller.activo else f"{taller.nombre} (inactivo)"
        _texto(hoja.cell(row=fila, column=1), nombre)
        _texto(hoja.cell(row=fila, column=2), taller.numero_taller)

        valores = (
            (3, taller.unidades_devueltas),
            (4, taller.monto_a_pagar),
            (5, taller.unidades_merma),
            # Excel guarda los porcentajes como fracción: 0.034 se muestra
            # como 3,4%. Dividir acá deja la celda como un porcentaje de
            # verdad, que promedia y grafica bien si el usuario lo intenta.
            (6, taller.tasa_merma / 100),
            (7, taller.entregas),
            (8, taller.dias_promedio_entrega),
            (9, taller.unidades_sin_tarifa),
            (10, taller.unidades_pendientes),
        )
        for columna, valor in valores:
            celda = hoja.cell(row=fila, column=columna)
            # Sin entregas no hay promedio de días: se deja la celda VACÍA.
            # Un 0 ahí se leería como "entrega el mismo día", que es lo
            # contrario de lo que pasó.
            if valor is not None:
                celda.value = valor
            celda.number_format = _COLUMNAS[columna - 1][2]

        if taller.unidades_sin_tarifa > 0:
            hoja.cell(row=fila, column=9).font = Font(bold=True, color=ROJO)

    return fila


def _escribir_totales(hoja: Worksheet, ranking: RankingTalleres, ultima_fila: int) -> int:
    fila = ultima_fila + 1
    relleno = PatternFill("solid", fgColor=GRIS_FONDO)
    borde_superior = Border(top=Side(style="thin", color=AZUL))

    _texto(hoja.cell(row=fila, column=1), "TOTAL")
    totales = {
        3: ranking.total_unidades_devueltas,
        4: ranking.total_a_pagar,
        5: ranking.total_unidades_merma,
        9: ranking.total_unidades_sin_tarifa,
    }
    # Las demás columnas no se totalizan a propósito: sumar días promedio o
    # porcentajes de merma daría un número sin significado.
    for columna in range(1, len(_COLUMNAS) + 1):
        celda = hoja.cell(row=fila, column=columna)
        if columna in totales:
            celda.value = totales[columna]
            celda.number_format = _COLUMNAS[columna - 1][2]
        celda.font = Font(bold=True)
        celda.fill = relleno
        celda.border = borde_superior
    return fila


def _escribir_notas(hoja: Worksheet, ranking: RankingTalleres, fila_total: int) -> None:
    fila = fila_total + 2

    if ranking.total_unidades_sin_tarifa > 0:
        aviso = hoja.cell(row=fila, column=1)
        _texto(
            aviso,
            f"ATENCIÓN: {ranking.total_unidades_sin_tarifa} unidades entregadas no tienen "
            f"precio cargado. El total a pagar está incompleto.",
        )
        aviso.font = Font(bold=True, color=ROJO)
        fila += 2

    nota = hoja.cell(row=fila, column=1)
    _texto(
        nota,
        "«Pendiente hoy» es lo que el taller tiene en su poder a la fecha de generación, "
        "no del periodo: parte de esas unidades se le asignaron antes del corte.",
    )
    nota.font = Font(size=9, italic=True)


def ranking_a_excel(ranking: RankingTalleres, empresa_nombre: str) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Talleres"

    _escribir_cabecera(hoja, ranking, empresa_nombre)
    ultima = _escribir_filas(hoja, ranking)
    fila_total = _escribir_totales(hoja, ranking, ultima)
    _escribir_notas(hoja, ranking, fila_total)

    # Con muchos talleres, al bajar se pierde de vista qué columna es cuál.
    hoja.freeze_panes = hoja.cell(row=_FILA_ENCABEZADO + 1, column=1)

    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def nombre_de_archivo(ranking: RankingTalleres) -> str:
    return f"talleres-{ranking.desde.isoformat()}-a-{ranking.hasta.isoformat()}.xlsx"
